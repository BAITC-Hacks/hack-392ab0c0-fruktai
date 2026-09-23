"""Export only persisted, computed recommendation rows; no automatic 1C writes."""

import csv
from io import BytesIO, StringIO
import json
from pathlib import Path

from database.repository import connect
from .imports import ImportProblem

EXPORT_HEADERS = [
    "sku",
    "name",
    "supplier_id",
    "supplier_name",
    "recommended_qty",
    "urgency",
    "reasons",
]


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_run(database_path, run_id, format, skus=None):
    if format not in {"csv", "xlsx"}:
        raise ImportProblem("Формат экспорта: csv или xlsx")
    with connect(database_path) as db:
        run = db.execute("SELECT 1 FROM calculation_runs WHERE run_id=?", (run_id,)).fetchone()
        if run is None:
            raise ImportProblem("Расчёт не найден", 404)
        rows = [
            dict(row)
            for row in db.execute(
                "SELECT * FROM recommendations WHERE run_id=? ORDER BY supplier_id, sku", (run_id,)
            )
        ]
    if skus:
        selected = set(skus)
        if selected - {row["sku"] for row in rows}:
            raise ImportProblem("В выбранном расчёте нет указанного артикула", 404)
        rows = [row for row in rows if row["sku"] in selected]
    values = []
    for row in rows:
        row["reasons"] = " | ".join(json.loads(row["reasons_json"]))
        values.append([safe_cell(row[key]) for key in EXPORT_HEADERS])
    if format == "csv":
        stream = StringIO(newline="")
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(EXPORT_HEADERS)
        writer.writerows(values)
        return stream.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "supplier_orders"
    sheet.append(EXPORT_HEADERS)
    for row in values:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def template_workbook():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    root = Path(__file__).resolve().parents[1] / "data/demo"
    for name in ("products", "suppliers", "sales", "inventory", "stockouts", "in_transit"):
        with (root / f"{name}.csv").open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            headers = list(reader.fieldnames)
            if name == "sales":
                headers += ["customer_id", "price", "warehouse_id"]
            elif name in {"inventory", "stockouts", "in_transit"}:
                headers += ["warehouse_id"]
            elif name == "products":
                headers += ["category", "unit"]
            sheet = workbook.create_sheet(name)
            sheet.append(headers)
            for row in reader:
                extra = {
                    "customer_id": "ANON-001",
                    "price": "100",
                    "warehouse_id": "WH-01",
                    "category": "Учебная категория",
                    "unit": "ед.",
                }
                sheet.append(
                    [safe_cell(row.get(header, extra.get(header, ""))) for header in headers]
                )
            sheet.freeze_panes = "A2"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()
