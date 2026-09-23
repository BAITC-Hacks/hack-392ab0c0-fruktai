"""Readers for CSV/TSV, JSON, Excel and text-table PDF; never an LLM."""

import csv
import json
import zipfile
from io import BytesIO, StringIO
from pathlib import Path
from .normalization import ImportProblem, MAX_ROWS, REQUIRED, table_name, normalize


def parse_file(filename, content, *, warnings=None):
    """Return normalized tables. Every sheet/row is validated, never skipped silently."""
    suffix = Path(filename).suffix.lower()
    tables = {}

    def add(name, matrix):
        key = table_name(name)
        if key in tables:
            raise ImportProblem(f"Повторная таблица {key}")
        tables[key] = normalize(key, matrix)

    try:
        if suffix in {".csv", ".tsv"}:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("cp1251")
            delimiter = (
                "\t"
                if suffix == ".tsv"
                else csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
            )
            matrix = list(csv.reader(StringIO(text), delimiter=delimiter))
            add(Path(filename).stem, matrix)
        elif suffix == ".json":
            value = json.loads(content)
            mapping = {Path(filename).stem: value} if isinstance(value, list) else value
            if not isinstance(mapping, dict):
                raise ImportProblem("JSON должен содержать таблицы или массив строк")
            for key, rows in mapping.items():
                name = table_name(key)
                if not isinstance(rows, list) or any(not isinstance(r, dict) for r in rows):
                    raise ImportProblem(f"{name}: ожидается массив объектов")
                headers = list(rows[0]) if rows else list(REQUIRED[name])
                if any(set(r) != set(headers) for r in rows):
                    raise ImportProblem(f"{name}: строки JSON имеют разные колонки")
                add(name, [headers] + [[r[h] for h in headers] for r in rows])
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            with zipfile.ZipFile(BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 60 * 1024 * 1024:
                    raise ImportProblem("Распакованная книга слишком велика", 413)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                if "Upload Data" in workbook.sheetnames:
                    from .logistiq import parse_logistiq

                    return parse_logistiq(workbook, read_excel_sheet, warnings)
                for sheet in workbook:
                    add(sheet.title, read_excel_sheet(sheet))
            finally:
                workbook.close()
        elif suffix == ".xls":
            import xlrd

            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                for sheet in workbook.sheets():
                    if sheet.nrows > MAX_ROWS + 1 or sheet.ncols > 32:
                        raise ImportProblem("Слишком большой лист Excel", 413)
                    matrix = []
                    for i in range(sheet.nrows):
                        values = []
                        for c in sheet.row(i):
                            if c.ctype == xlrd.XL_CELL_ERROR:
                                raise ImportProblem(f"{sheet.name}: ошибка в ячейке XLS")
                            values.append(
                                xlrd.xldate.xldate_as_datetime(c.value, workbook.datemode)
                                if c.ctype == xlrd.XL_CELL_DATE
                                else c.value
                            )
                        matrix.append(values)
                    add(sheet.name, matrix)
            finally:
                workbook.release_resources()
        elif suffix == ".pdf":
            import pdfplumber

            matrix = []
            with pdfplumber.open(BytesIO(content)) as pdf:
                if len(pdf.pages) > 50:
                    raise ImportProblem("PDF: максимум 50 страниц", 413)
                for page in pdf.pages:
                    if not (page.extract_text() or "").strip():
                        raise ImportProblem(
                            "PDF содержит скан/страницу без текста. Нужен OCR или выгрузка XLSX/CSV."
                        )
                    extracted = page.extract_tables()
                    if not extracted:
                        raise ImportProblem(
                            "В PDF не найдены таблицы с границами. Экспортируйте XLSX/CSV."
                        )
                    for table in extracted:
                        matrix.extend(table)
            add(Path(filename).stem, matrix)
        else:
            raise ImportProblem("Поддерживаются CSV, TSV, XLSX, XLS, JSON и текстовые PDF", 415)
    except ImportProblem:
        raise
    except Exception as error:
        raise ImportProblem(
            f"Не удалось прочитать {suffix} файл: проверьте формат, пароль и структуру таблицы"
        ) from error
    return tables


def read_excel_sheet(sheet):
    """Enforce limits while streaming, including XLSX without dimension metadata."""
    if (sheet.max_row or 0) > MAX_ROWS + 3 or (sheet.max_column or 0) > 32:
        raise ImportProblem("Слишком большой лист Excel", 413)
    matrix = []
    for row in sheet.iter_rows():
        if len(matrix) >= MAX_ROWS + 3 or len(row) > 32:
            raise ImportProblem("Слишком большой лист Excel", 413)
        if any(c.data_type in {"f", "e"} for c in row):
            raise ImportProblem(f"{sheet.title}: замените формулы/ошибки значениями")
        matrix.append([c.value for c in row])
    return matrix
