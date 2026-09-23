"""Synthetic LOGISTIQ regression fixtures; no user workbook checked into Git."""

from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from agent import WorkflowService
from backend.file_exchange.logistiq import SNAPSHOT_COLUMNS, UPLOAD_COLUMNS
from backend.imports import ImportProblem, parse_file
from backend.main import create_app


def workbook_bytes(*, mutate=None, reference=99999):
    history = []
    for index in range(28):
        history.append(
            dict(
                zip(
                    UPLOAD_COLUMNS,
                    [
                        date(2026, 9, 1) + timedelta(days=index),
                        "WH-1",
                        "001",
                        "Part",
                        "Parts",
                        "Supplier",
                        0 if index == 10 else 10,
                        reference,
                        reference,
                        0 if index == 10 else 99,
                        0,
                        None,
                        7,
                        5,
                        10,
                        25,
                        0.95,
                        False,
                        index == 10,
                        False,
                        "OK",
                    ],
                )
            )
        )
    snapshot = [
        dict(
            zip(
                SNAPSHOT_COLUMNS,
                [
                    "001",
                    "Part",
                    "Parts",
                    "Supplier",
                    20,
                    5,
                    date(2026, 10, 1),
                    7,
                    5,
                    10,
                    0,
                    0.95,
                    reference,
                    reference,
                    reference,
                    reference,
                    reference,
                    reference,
                    "Low",
                    "Scenario",
                    "Reference only",
                ],
            )
        )
    ]
    if mutate:
        mutate(history, snapshot)
    # write_only produces worksheets without optional dimension metadata.
    book = Workbook(write_only=True)
    for name in ("Expected Results", "Data Dictionary"):
        sheet = book.create_sheet(name)
        sheet.append(["Reference only", reference])
    for name, headers, rows in (
        ("Upload Data", UPLOAD_COLUMNS, history),
        ("Current Snapshot", SNAPSHOT_COLUMNS, snapshot),
    ):
        sheet = book.create_sheet(name)
        sheet.append([name])
        sheet.append(["Synthetic test input"])
        sheet.append(headers)
        for row in rows:
            sheet.append([row[h] for h in headers])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_dimensionless_workbook_and_explicit_snapshot_precedence():
    warnings = []
    result = parse_file("input.xlsx", workbook_bytes(), warnings=warnings)
    assert len(result["sales"]) == 28
    assert result["inventory"] == [{"sku": "001", "warehouse_id": "WH-1", "on_hand": "20"}]
    assert result["in_transit"][0]["quantity"] == "5"
    assert result["stockouts"][0]["start_date"] == "2026-09-11"
    assert all("price" not in row for row in result["sales"])
    assert any("Current Snapshot" in message for message in warnings)
    assert any("MOQ" in message for message in warnings)
    assert result == parse_file("renamed.xlsx", workbook_bytes(reference=123456))


@pytest.mark.parametrize(
    "mutate",
    [
        lambda h, s: s[0].update({"На складе": -1}),
        lambda h, s: s[0].update({"ETA": None}),
        lambda h, s: s.append(s[0].copy()),
        lambda h, s: h.append(h[0].copy()),
        lambda h, s: h[0].update(warehouse="WH-2"),
        lambda h, s: h[0].update(stockout_flag="maybe"),
        lambda h, s: h[0].update(sales_qty="=1+1"),
        lambda h, s: h[0].update(product_name="Different"),
        lambda h, s: h[0].update(sku="Unknown"),
        lambda h, s: h[10].update(stock_on_hand=1),
    ],
)
def test_invalid_workbook_is_rejected(mutate):
    with pytest.raises(ImportProblem):
        parse_file("input.xlsx", workbook_bytes(mutate=mutate))


def test_real_api_import_calculate_override_export(tmp_path):
    service = WorkflowService(Path("data").resolve(), tmp_path / "db.sqlite", tmp_path / "runs")
    with TestClient(create_app(service)) as client:
        response = client.post(
            "/api/v1/datasets/import",
            data={"anonymized": "true"},
            files={"files": ("input.xlsx", workbook_bytes())},
        )
        assert response.status_code == 201, response.text
        meta = response.json()
        assert meta["warehouse_id"] == "WH-1"
        assert meta["counts"]["sales"] == 28
        response = client.post("/api/v1/recalculate", json={"dataset": meta["dataset"]})
        assert response.status_code == 200, response.text
        result = response.json()
        item = result["recommendations"][0]
        assert item["on_hand"] == 20
        assert item["in_transit"] == 5
        assert item["recommended_qty"] > 0
        assert item["stockout_compensation"] > 0
        changed = client.post(
            "/api/v1/recalculate",
            json={
                "dataset": meta["dataset"],
                "overrides": [{"sku": "001", "on_hand": 10000}],
            },
        )
        assert changed.status_code == 200
        assert changed.json()["recommendations"][0]["recommended_qty"] == 0
        assert client.get(f"/api/v1/runs/{result['run_id']}/export?format=xlsx").status_code == 200


def test_canonical_dimensionless_xlsx_still_supported():
    book = Workbook(write_only=True)
    sheet = book.create_sheet("inventory")
    sheet.append(["sku", "on_hand"])
    sheet.append(["001", 20])
    output = BytesIO()
    book.save(output)
    book.close()
    assert parse_file("inventory.xlsx", output.getvalue()) == {
        "inventory": [{"sku": "001", "on_hand": "20"}],
    }
