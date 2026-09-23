import csv
from io import BytesIO, StringIO
import json
from pathlib import Path
import sqlite3

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from agent import WorkflowService
from backend.main import create_app
from backend.imports import parse_file, ImportProblem
from scripts.validate_api_contracts import load_schemas, validate

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def client(tmp_path):
    service = WorkflowService(ROOT / "data", tmp_path / "db.sqlite", tmp_path / "runs")
    with TestClient(create_app(service)) as client:
        yield client


def import_json(client, tables, **form):
    return client.post(
        "/api/v1/datasets/import",
        data={"anonymized": "true", **form},
        files={"files": ("input.json", json.dumps(tables), "application/json")},
    )


def source_tables():
    tables = {}
    for path in (ROOT / "data/demo").glob("*.csv"):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            tables[path.stem] = list(csv.DictReader(stream))
    return tables


def test_template_to_calculation_to_both_exports(client):
    template = client.get("/api/v1/datasets/template.xlsx")
    assert template.status_code == 200
    response = client.post(
        "/api/v1/datasets/import",
        data={"anonymized": "true"},
        files={"files": ("input.xlsx", template.content)},
    )
    assert response.status_code == 201, response.text
    meta = response.json()
    schema = load_schemas()["import.response.schema.json"]
    validate(meta, schema, schema)
    assert meta["counts"]["sales"] == 336
    assert meta["warehouse_id"] == "WH-01"
    assert not meta["warnings"]
    assert "customer_id" not in json.dumps(meta["preview"])
    assert client.get("/api/v1/datasets/" + meta["dataset"]).json() == meta
    result = client.post("/api/v1/recalculate", json={"dataset": meta["dataset"]})
    assert result.status_code == 200, result.text
    body = result.json()
    assert body["summary"]["estimated_stockout_items"] > 0
    assert body["summary"]["anomalies_removed"] > 0
    run = body["run_id"]
    item = body["recommendations"][0]
    csv_export = client.get(f"/api/v1/runs/{run}/export?format=csv&sku={item['sku']}")
    rows = list(csv.DictReader(StringIO(csv_export.content.decode("utf-8-sig")), delimiter=";"))
    assert len(rows) == 1
    assert int(rows[0]["recommended_qty"]) == item["recommended_qty"]
    xlsx = client.get(f"/api/v1/runs/{run}/export?format=xlsx")
    book = load_workbook(BytesIO(xlsx.content), read_only=True)
    assert book.active.max_row == 7
    book.close()
    # Later calculations for the same SKU cannot change this run's export/history.
    client.post(
        "/api/v1/recalculate",
        json={"dataset": meta["dataset"], "overrides": [{"sku": item["sku"], "on_hand": 10000}]},
    )
    detail = client.get(f"/api/v1/items/{item['sku']}?run_id={run}").json()
    assert detail["calculation"]["recommended_qty"] == item["recommended_qty"]
    assert (
        client.get(f"/api/v1/runs/{run}/export?format=csv&sku={item['sku']}").content
        == csv_export.content
    )


def test_1c_material_statement_and_warehouse(client):
    tables = source_tables()
    for name in ("sales", "inventory", "stockouts", "in_transit"):
        tables[name] = [dict(row, warehouse_id="WH-A") for row in tables[name]]
    tables["material_statement"] = tables.pop("inventory")
    tables["material_statement"].append(
        {"sku": "STABLE-001", "on_hand": "900", "warehouse_id": "WH-B"}
    )
    missing = import_json(client, tables)
    assert missing.status_code == 422
    response = import_json(client, tables, warehouse_id="WH-A")
    assert response.status_code == 201, response.text
    assert response.json()["source"] == "1c_file_exchange"
    tables["inventory"] = [dict(row, on_hand="999") for row in tables["material_statement"]]
    assert import_json(client, tables, warehouse_id="WH-A").status_code == 422


def test_atomic_invalid_import_and_anonymization(client, tmp_path):
    tables = source_tables()
    tables["inventory"][0]["on_hand"] = "-1"
    assert import_json(client, tables).status_code == 422
    with sqlite3.connect(tmp_path / "db.sqlite") as db:
        assert db.execute("SELECT count(*) FROM source_imports").fetchone()[0] == 0
    response = client.post(
        "/api/v1/datasets/import",
        data={"anonymized": "false"},
        files={"files": ("input.json", json.dumps(source_tables()))},
    )
    assert response.status_code == 422
    assert import_json(client, source_tables()).status_code == 201


def test_unsupported_missing_tables_and_formula(client):
    assert (
        client.post(
            "/api/v1/datasets/import",
            data={"anonymized": "true"},
            files={"files": ("data.exe", b"bad")},
        ).status_code
        == 415
    )
    assert import_json(client, {"sales": []}).status_code == 422
    book = load_workbook(BytesIO(client.get("/api/v1/datasets/template.xlsx").content))
    book["sales"]["C2"] = "=1+1"
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    assert (
        client.post(
            "/api/v1/datasets/import",
            data={"anonymized": "true"},
            files={"files": ("input.xlsx", buffer.getvalue())},
        ).status_code
        == 422
    )


def test_russian_cp1251_csv_and_tsv():
    result = parse_file(
        "Продажи.csv", "Дата;Артикул;Количество;Цена\n22.09.2026;0001;12,5;20,2\n".encode("cp1251")
    )
    assert result["sales"][0] == {
        "date": "2026-09-22",
        "sku": "0001",
        "units": "12.5",
        "price": "20.2",
    }
    assert parse_file("inventory.tsv", b"sku\ton_hand\n001\t25")["inventory"][0]["sku"] == "001"


def simple_pdf(text=True):
    # A real minimal PDF with table borders. No external PDF generation dependency.
    commands = (
        [
            "BT /F1 10 Tf 50 710 Td (sku) Tj 150 0 Td (on_hand) Tj ET",
            "BT /F1 10 Tf 50 680 Td (001) Tj 150 0 Td (25) Tj ET",
        ]
        if text
        else []
    )
    for y in (730, 700, 670):
        commands.append(f"40 {y} m 300 {y} l S")
    for x in (40, 190, 300):
        commands.append(f"{x} 670 m {x} 730 l S")
    stream = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, value in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{i} 0 obj\n".encode() + value + b"\nendobj\n")
    position = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{position}\n%%EOF".encode()
    )
    return bytes(pdf)


def test_real_pdf_tables_and_scan_rejection():
    assert parse_file("inventory.pdf", simple_pdf())["inventory"] == [
        {"sku": "001", "on_hand": "25"}
    ]
    with pytest.raises(ImportProblem, match="скан"):
        parse_file("inventory.pdf", simple_pdf(False))


def test_import_preserves_price_and_customer_snapshot(client, tmp_path):
    tables = source_tables()
    for row in tables["sales"]:
        row["customer_id"] = "ANON-001"
        row["price"] = "120.5"
    response = import_json(client, tables)
    assert response.status_code == 201
    with sqlite3.connect(tmp_path / "db.sqlite") as db:
        snapshot = json.loads(db.execute("SELECT tables_json FROM source_imports").fetchone()[0])
        assert snapshot["sales"][0]["price"] == "120.5"
        assert snapshot["sales"][0]["customer_id"] == "ANON-001"


def test_real_legacy_xls():
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("inventory")
    for i, row in enumerate((("sku", "on_hand"), ("0001", 25))):
        for j, value in enumerate(row):
            sheet.write(i, j, value)
    output = BytesIO()
    book.save(output)
    assert parse_file("input.xls", output.getvalue())["inventory"] == [
        {"sku": "0001", "on_hand": "25"}
    ]


def test_client_spike_preserves_other_sales_and_price_does_not_set_quantity(client):
    from datetime import date, timedelta

    tables = {
        "products": [{"sku": "001", "name": "Товар", "supplier_id": "S1", "active": "true"}],
        "suppliers": [{"supplier_id": "S1", "supplier_name": "Поставщик", "lead_time_days": "7"}],
        "inventory": [{"sku": "001", "on_hand": "0"}],
        "stockouts": [],
        "in_transit": [],
        "sales": [
            {
                "sku": "001",
                "date": (date(2026, 1, 1) + timedelta(days=i)).isoformat(),
                "units": "10",
                "customer_id": "REGULAR",
                "price": "5",
            }
            for i in range(28)
        ],
    }
    tables["sales"].append(dict(tables["sales"][-1], units="1000", customer_id="ONE-OFF"))
    dataset = import_json(client, tables).json()["dataset"]
    result = client.post("/api/v1/recalculate", json={"dataset": dataset}).json()
    item = result["recommendations"][0]
    assert item["avg_daily_demand"] == 10
    assert item["outlier_units_removed"] == 1000
    assert item["recommended_qty"] == 140
    detail = client.get("/api/v1/items/001?run_id=" + result["run_id"]).json()
    assert detail["history"][-1]["units"] == 1010
    assert detail["history"][-1]["is_outlier"]
    for row in tables["sales"]:
        row["price"] = "9000"
    other = import_json(client, tables).json()["dataset"]
    unchanged = client.post("/api/v1/recalculate", json={"dataset": other}).json()
    assert unchanged["recommendations"][0]["recommended_qty"] == 140
    changed = client.post(
        "/api/v1/recalculate",
        json={"dataset": dataset, "overrides": [{"sku": "001", "in_transit": 200}]},
    ).json()
    assert changed["recommendations"][0]["recommended_qty"] == 0


def test_export_formula_injection_and_invalid_selection(client):
    tables = source_tables()
    tables["products"][0]["name"] = "=HYPERLINK(1)"
    dataset = import_json(client, tables).json()["dataset"]
    result = client.post("/api/v1/recalculate", json={"dataset": dataset}).json()
    run = result["run_id"]
    response = client.get(f"/api/v1/runs/{run}/export?format=csv")
    assert "'=HYPERLINK(1)" in response.content.decode("utf-8-sig")
    assert client.get(f"/api/v1/runs/{run}/export?format=xlsx&sku=missing").status_code == 404
    assert client.get(f"/api/v1/runs/{run}/export?format=pdf").status_code == 422


def test_oversized_stockout_period_rejected_before_expansion(client):
    tables = source_tables()
    tables["stockouts"] = [
        {"sku": tables["products"][0]["sku"], "start_date": "0001-01-01", "end_date": "9999-12-31"}
    ]
    assert import_json(client, tables).status_code == 413
